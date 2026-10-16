"""Editable poster renderer with bounded glyph-targeted inset recovery.

Default rendering is identical to the shared renderer. Only actually failed
text targets receive larger insets. Allocation/wrapping in layout.py matches
these insets; shared glyph inspection is unchanged and always runs.
"""
from pathlib import Path
import os
import shutil
import subprocess
import uuid
from dagent.core.config import writable_path, dump_json
from dagent.core.layout.geometry import wrap_lines
from dagent.core.layout.render_checks import inspect_rendered_text

class PosterRenderer:
    def __init__(self, cfg, ledger, *, text_safety=None):
        self.cfg,self.ledger=cfg,ledger
        self.text_safety = text_safety or {}

    def render(self, layout, directory: Path):
        import fitz
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR, PP_ALIGN
        from dagent.core.pptx_utils import create_poster, add_blank_slide, add_textbox, add_image, save_presentation
        directory=writable_path(directory)
        directory.mkdir(parents=True,exist_ok=True)
        executable=shutil.which('soffice') or shutil.which('libreoffice')
        if not executable:
            raise RuntimeError('LibreOffice/soffice is required for actual PPTX rendering (no synthetic preview fallback)')
        prs=create_poster(layout.width,layout.height)
        slide=add_blank_slide(prs)
        for e in layout.elements:
            if e.kind=='figure':
                add_image(slide,e.id,e.x,e.y,e.width,e.height,e.asset_path)
                continue
            shape=add_textbox(slide,e.id,e.x,e.y,e.width,e.height,text='',font_size=e.font_pt,
                              bold=e.bold,fill_color=tuple(e.fill) if e.fill else None,font_name='Arial')
            tf=shape.text_frame
            tf.clear()
            tf.word_wrap=False  # explicit metric-based line breaks, stable editable text
            tf.auto_size=MSO_AUTO_SIZE.NONE
            tf.vertical_anchor=MSO_ANCHOR.MIDDLE if self.cfg.design in ('balanced', 'portfolio', 'editorial') and e.kind == 'title' else MSO_ANCHOR.TOP
            repairs = min(2, max(0, self.text_safety.get(e.id, 0)))
            tf.margin_top=Inches(.08 + .10 * repairs)
            tf.margin_bottom=Inches(.08)
            tf.margin_left=tf.margin_right=Inches(.08 + .06 * repairs)
            for i,line in enumerate(wrap_lines(e.text,e.width - .12 * repairs,e.font_pt,e.bold)):
                p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
                p.text=line
                if self.cfg.design in ('balanced', 'portfolio', 'editorial') and e.kind in ('title', 'authors'):
                    p.alignment=PP_ALIGN.CENTER
                p.font.name='Arial'
                p.font.size=Pt(e.font_pt)
                p.font.bold=e.bold
                p.font.color.rgb=RGBColor(*e.color)
                p.space_before=p.space_after=Pt(0)
                # Use the same point-based line height as geometry.text_height.
                # A float multiplies the font's *natural* height in LibreOffice.
                p.line_spacing=Pt(e.font_pt * e.line_spacing)
        pptx=writable_path(directory/'poster.pptx')
        save_presentation(prs,str(pptx))
        profile=writable_path(directory/('lo_profile_'+uuid.uuid4().hex))
        temp=writable_path(directory/'tmp')
        temp.mkdir(exist_ok=True)
        runtime=writable_path(temp/'runtime')
        runtime.mkdir(exist_ok=True,mode=0o700)
        env=dict(os.environ,TMPDIR=str(temp),HOME=str(temp),XDG_CACHE_HOME=str(temp/'cache'),
                 XDG_CONFIG_HOME=str(temp/'config'),XDG_RUNTIME_DIR=str(runtime),
                 DCONF_BACKEND='memory',SAL_USE_VCLPLUGIN='svp')
        result=subprocess.run([executable,f'-env:UserInstallation={profile.as_uri()}',
            '--headless','--convert-to','pdf','--outdir',str(directory),str(pptx)],
            capture_output=True,text=True,timeout=self.cfg.renderer_timeout_seconds,env=env,cwd=directory)
        self.ledger.renderer_calls+=1
        self.ledger.save()
        pdf=directory/'poster.pdf'
        if result.returncode!=0 or not pdf.exists():
            raise RuntimeError(f'LibreOffice conversion failed: {result.stderr[-700:]} {result.stdout[-700:]}')
        png=writable_path(directory/'poster.png')
        with fitz.open(pdf) as doc:
            if len(doc)!=1:
                raise RuntimeError('Expected a single-slide poster PDF')
            page=doc[0]
            page.get_pixmap(matrix=fitz.Matrix(self.cfg.render_dpi/72,self.cfg.render_dpi/72),alpha=False).save(png)
            render_defects = inspect_rendered_text(page, layout)
        dump_json(directory/'render_checks.json',dict(defects=render_defects,valid=not render_defects))
        dump_json(directory/'layout.json',layout.to_dict())
        return dict(pptx=pptx,png=png,pdf=pdf,defects=render_defects)
